from __future__ import annotations

import hashlib
import hmac
import json
import mimetypes
import secrets
import smtplib
import time
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from email.message import EmailMessage

from app.email_design import branded_email_html, attach_sophies_logo
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Form, HTTPException, Request
from fastapi.responses import (
    FileResponse,
    HTMLResponse,
    RedirectResponse,
    StreamingResponse,
)
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import select

from app.auth.service import new_csrf_token
from app.db.models import (
    StaffMember,
    StaffPayRate,
    StaffPaymentProfile,
    StaffPayrollSubmission,
)
from app.storage import download_to_temp


def _payroll_csrf(request: Request) -> str:
    token = new_csrf_token()
    request.session["payroll_admin_csrf_token"] = token
    return token


def _valid_payroll_csrf(
    request: Request,
    supplied: str,
) -> bool:
    expected = str(
        request.session.get(
            "payroll_admin_csrf_token",
            "",
        )
    )

    return bool(
        expected
        and hmac.compare_digest(
            expected,
            str(supplied or ""),
        )
    )


def _payload(row: StaffPayrollSubmission) -> dict:
    try:
        value = json.loads(row.payload_json or "{}")
    except Exception:
        value = {}

    if not isinstance(value, dict):
        value = {}

    if not isinstance(value.get("timesheet_entries"), list):
        value["timesheet_entries"] = []

    if not isinstance(value.get("expense_claims"), list):
        value["expense_claims"] = []

    return value


def _money(value: object) -> str:
    try:
        return f"{float(str(value or '0').replace(',', '')):.2f}"
    except (TypeError, ValueError):
        return "0.00"


def _minutes(start: object, finish: object, break_minutes: object = 0) -> int:
    def clock(value: object) -> int | None:
        text = str(value or "").strip()
        if not text or ":" not in text:
            return None
        try:
            hour, minute = text.split(":", 1)
            return int(hour) * 60 + int(minute[:2])
        except (TypeError, ValueError):
            return None

    start_value = clock(start)
    finish_value = clock(finish)

    if start_value is None or finish_value is None:
        return 0

    if finish_value < start_value:
        finish_value += 24 * 60

    try:
        break_value = max(0, int(break_minutes or 0))
    except (TypeError, ValueError):
        break_value = 0

    return max(0, finish_value - start_value - break_value)



def _contractor_statement_pdf(
    *,
    request: Request,
    submission: StaffPayrollSubmission,
    payload: dict,
) -> bytes:
    """
    Branded self-employed contractor payment statement.

    This is not a PAYE payslip.
    """

    arrangement = str(
        payload.get("payment_arrangement")
        or ""
    ).strip()

    if arrangement != "self_employed_contractor":
        raise HTTPException(
            409,
            "Contractor statements are only available for self-employed contractors.",
        )

    timesheets = [
        item
        for item in payload.get(
            "timesheet_entries",
            [],
        )
        if isinstance(item, dict)
    ]

    expenses = [
        item
        for item in payload.get(
            "expense_claims",
            [],
        )
        if isinstance(item, dict)
    ]

    paid_minutes = sum(
        _minutes(
            item.get("start_time"),
            item.get("finish_time"),
            item.get("break_minutes"),
        )
        for item in timesheets
    )

    hours = paid_minutes / 60

    hourly_rate = float(
        _money(
            payload.get("hourly_rate")
        )
    )

    gross_pay = (
        hours
        * hourly_rate
    )

    expense_total = sum(
        float(
            _money(
                item.get("amount")
            )
        )
        for item in expenses
    )

    total_due = (
        gross_pay
        + expense_total
    )

    week_end = (
        submission.week_start
        + timedelta(days=6)
    )

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=(
            f"Contractor Payment Statement - "
            f"{submission.staff_name}"
        ),
        author="Sophie's Photography",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "StatementTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#111111"),
        spaceAfter=4 * mm,
    )

    heading_style = ParagraphStyle(
        "StatementHeading",
        parent=styles["Heading2"],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#111111"),
        spaceBefore=4 * mm,
        spaceAfter=2 * mm,
    )

    normal_style = ParagraphStyle(
        "StatementNormal",
        parent=styles["BodyText"],
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#222222"),
    )

    small_style = ParagraphStyle(
        "StatementSmall",
        parent=normal_style,
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#555555"),
    )

    right_style = ParagraphStyle(
        "StatementRight",
        parent=normal_style,
        alignment=TA_RIGHT,
    )

    story = []

    logo_path = (
        Path(__file__).resolve().parent
        / "static"
        / "images"
        / "sophies-dashboard-logo.png"
    )

    if logo_path.exists():
        try:
            logo = Image(
                str(logo_path),
                width=45 * mm,
                height=18 * mm,
                kind="proportional",
            )

            story.append(logo)
            story.append(
                Spacer(
                    1,
                    4 * mm,
                )
            )
        except Exception:
            pass

    story.append(
        Paragraph(
            "CONTRACTOR PAYMENT STATEMENT",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "Sophie’s Photography",
            normal_style,
        )
    )

    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )

    statement_ref = (
        "CPS-"
        + str(
            submission.id
        ).split("-")[0].upper()
        + "-"
        + submission.week_start.strftime(
            "%Y%m%d"
        )
    )

    meta_data = [
        [
            Paragraph(
                "<b>Contractor</b>",
                normal_style,
            ),
            Paragraph(
                str(
                    submission.staff_name
                    or ""
                ),
                normal_style,
            ),
            Paragraph(
                "<b>Statement reference</b>",
                normal_style,
            ),
            Paragraph(
                statement_ref,
                right_style,
            ),
        ],
        [
            Paragraph(
                "<b>Email</b>",
                normal_style,
            ),
            Paragraph(
                str(
                    submission.staff_email
                    or ""
                ),
                normal_style,
            ),
            Paragraph(
                "<b>Payroll week</b>",
                normal_style,
            ),
            Paragraph(
                (
                    f"{submission.week_start.strftime('%d %b %Y')}"
                    " - "
                    f"{week_end.strftime('%d %b %Y')}"
                ),
                right_style,
            ),
        ],
    ]

    meta = Table(
        meta_data,
        colWidths=[
            28 * mm,
            55 * mm,
            32 * mm,
            52 * mm,
        ],
    )

    meta.setStyle(
        TableStyle(
            [
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
            ]
        )
    )

    story.append(meta)

    story.append(
        Paragraph(
            "Work Detail",
            heading_style,
        )
    )

    work_rows = [
        [
            "Date",
            "Start",
            "Finish",
            "Break",
            "Hours",
            "Description",
        ]
    ]

    for item in timesheets:

        minutes = _minutes(
            item.get("start_time"),
            item.get("finish_time"),
            item.get("break_minutes"),
        )

        description = str(
            item.get("notes")
            or item.get("work_type")
            or item.get("event_name")
            or "Work"
        )

        work_rows.append(
            [
                str(
                    item.get(
                        "work_date",
                        "",
                    )
                ),
                str(
                    item.get(
                        "start_time",
                        "",
                    )
                ),
                str(
                    item.get(
                        "finish_time",
                        "",
                    )
                ),
                str(
                    item.get(
                        "break_minutes",
                        0,
                    )
                )
                + " min",
                f"{minutes / 60:.2f}",
                Paragraph(
                    description,
                    small_style,
                ),
            ]
        )

    if len(work_rows) == 1:
        work_rows.append(
            [
                "—",
                "—",
                "—",
                "—",
                "0.00",
                "No work entries",
            ]
        )

    work_table = Table(
        work_rows,
        repeatRows=1,
        colWidths=[
            23 * mm,
            17 * mm,
            17 * mm,
            18 * mm,
            16 * mm,
            76 * mm,
        ],
    )

    work_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#111111"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.5,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.HexColor(
                        "#cccccc"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
            ]
        )
    )

    story.append(work_table)

    if expenses:
        story.append(
            Paragraph(
                "Approved Expenses",
                heading_style,
            )
        )

        expense_rows = [
            [
                "Date",
                "Description",
                "Amount",
            ]
        ]

        for item in expenses:
            expense_rows.append(
                [
                    str(
                        item.get(
                            "expense_date",
                            "",
                        )
                    ),
                    Paragraph(
                        str(
                            item.get(
                                "description"
                            )
                            or item.get(
                                "expense_type"
                            )
                            or "Expense"
                        ),
                        small_style,
                    ),
                    "£"
                    + _money(
                        item.get(
                            "amount"
                        )
                    ),
                ]
            )

        expense_table = Table(
            expense_rows,
            repeatRows=1,
            colWidths=[
                30 * mm,
                110 * mm,
                27 * mm,
            ],
        )

        expense_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#111111"
                        ),
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7.5,
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.3,
                        colors.HexColor(
                            "#cccccc"
                        ),
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                ]
            )
        )

        story.append(
            expense_table
        )

    story.append(
        Paragraph(
            "Payment Summary",
            heading_style,
        )
    )

    summary_rows = [
        [
            "Total hours",
            f"{hours:.2f}",
        ],
        [
            "Agreed hourly rate",
            f"£{hourly_rate:.2f}",
        ],
        [
            "Work value",
            f"£{gross_pay:.2f}",
        ],
        [
            "Approved expenses",
            f"£{expense_total:.2f}",
        ],
        [
            "Total payment",
            f"£{total_due:.2f}",
        ],
    ]

    if submission.paid_at:
        summary_rows.extend(
            [
                [
                    "Payment status",
                    "PAID",
                ],
                [
                    "Payment reference",
                    str(
                        submission.payment_reference
                        or ""
                    ),
                ],
            ]
        )
    else:
        summary_rows.append(
            [
                "Payment status",
                (
                    "APPROVED - AWAITING PAYMENT"
                    if submission.approved_at
                    else "NOT YET APPROVED"
                ),
            ]
        )

    summary_table = Table(
        summary_rows,
        colWidths=[
            95 * mm,
            72 * mm,
        ],
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.HexColor(
                        "#cccccc"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 4),
                    (-1, 4),
                    "Helvetica-Bold",
                ),
                (
                    "BACKGROUND",
                    (0, 4),
                    (-1, 4),
                    colors.HexColor(
                        "#f0f0f0"
                    ),
                ),
                (
                    "ALIGN",
                    (1, 0),
                    (1, -1),
                    "RIGHT",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8.5,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(
            1,
            6 * mm,
        )
    )

    story.append(
        Paragraph(
            (
                "This document is a self-employed contractor "
                "payment statement and is not a PAYE payslip. "
                "It records work and payments made by Sophie’s "
                "Photography. The contractor remains responsible "
                "for their own tax affairs and for declaring "
                "income to HMRC where required."
            ),
            small_style,
        )
    )

    story.append(
        Spacer(
            1,
            3 * mm,
        )
    )

    story.append(
        Paragraph(
            (
                "Generated securely by Stuphie Ltd · "
                + datetime.now().astimezone().strftime(
                    "%d %B %Y %H:%M"
                )
            ),
            small_style,
        )
    )

    doc.build(story)

    return buffer.getvalue()


def _expense_receipt(expense: dict) -> str:
    for key in (
        "receipt_storage_path",
        "receipt_storage_location",
        "receipt_location",
        "receipt_path",
    ):
        value = str(expense.get(key) or "").strip()
        if value:
            return value
    return ""


def _expense_filename(expense: dict, location: str) -> str:
    for key in ("receipt_filename", "receipt_name", "filename"):
        value = str(expense.get(key) or "").strip()
        if value:
            return Path(value).name

    if location:
        return Path(location.split("?", 1)[0]).name or "receipt"

    return "receipt"


PAYROLL_OTP_MAX_ATTEMPTS = 5
PAYROLL_OTP_RESEND_SECONDS = 60


def _payroll_otp_digest(secret_key: str, code: str) -> str:
    return hmac.new(
        str(secret_key).encode("utf-8"),
        str(code).encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()


def _payroll_identity(request: Request) -> str:
    return str(request.session.get("staff_email") or "").strip().lower()


def _require_staff_super_admin(request: Request) -> str:
    if request.session.get("staff_role") != "super_admin":
        raise HTTPException(403, "Super-admin access required.")

    email = _payroll_identity(request)

    if not email:
        raise HTTPException(403, "Authenticated staff email required.")

    return email


def _payroll_email_allowed(request: Request, email: str) -> bool:
    settings = request.app.state.settings

    return email in settings.payroll_otp_allowed_email_set


def _clear_pending_payroll_otp(request: Request) -> None:
    for key in (
        "payroll_otp_email",
        "payroll_otp_digest",
        "payroll_otp_expires_at",
        "payroll_otp_attempts",
        "payroll_otp_last_sent_at",
    ):
        request.session.pop(key, None)


def _clear_payroll_unlock(request: Request) -> None:
    for key in (
        "payroll_otp_verified_at",
        "payroll_otp_last_activity",
    ):
        request.session.pop(key, None)


def _require_super_admin(request: Request) -> None:
    """
    Payroll Admin requires both Super Admin authentication and
    a recent Payroll email-OTP verification.
    """
    email = _require_staff_super_admin(request)

    if not _payroll_email_allowed(request, email):
        _clear_pending_payroll_otp(request)
        _clear_payroll_unlock(request)
        raise HTTPException(
            403,
            "This Super Admin account is not authorised for Payroll Admin.",
        )

    now = time.time()

    verified_at = float(
        request.session.get("payroll_otp_verified_at") or 0
    )
    last_activity = float(
        request.session.get("payroll_otp_last_activity") or 0
    )

    settings = request.app.state.settings
    unlock_seconds = int(settings.payroll_otp_unlock_seconds)

    if (
        verified_at <= 0
        or last_activity <= 0
        or now - last_activity > unlock_seconds
    ):
        _clear_payroll_unlock(request)

        raise HTTPException(
            status_code=303,
            headers={"Location": "/staff/payroll/unlock"},
        )

    request.session["payroll_otp_last_activity"] = now


def _send_payroll_otp(settings, recipient: str, code: str) -> None:
    if not settings.smtp_host:
        raise RuntimeError("Payroll authentication email is unavailable.")

    message = EmailMessage()
    message["From"] = (
        f"{settings.smtp_from_name} "
        f"<{settings.smtp_from_email}>"
    )
    message["To"] = recipient
    message["Subject"] = "Your Stuphie Payroll security code"

    message.set_content(
        "Your Stuphie Payroll security code is:\n\n"
        f"{code}\n\n"
        "This code expires in 10 minutes and can only be used once.\n"
        "If you did not request this code, you can ignore this email.\n\n"
        "Sophie’s Photography"
    )

    message.add_alternative(
        branded_email_html(
            "Payroll security code",
            f"Your Stuphie Payroll security code is:\n\n{code}\n\nThis code expires in 10 minutes and can only be used once. If you did not request it, you can ignore this email.",
            eyebrow="SOPHIE’S PHOTOGRAPHY · STAFF",
            include_marketing_controls=False,
        ),
        subtype="html",
    )
    attach_sophies_logo(message)

    with smtplib.SMTP(
        settings.smtp_host,
        settings.smtp_port,
        timeout=20,
    ) as smtp:
        if settings.smtp_use_tls:
            smtp.starttls()

        if settings.smtp_username:
            smtp.login(
                settings.smtp_username,
                settings.smtp_password,
            )

        smtp.send_message(message)


def _month_value(value: str | None) -> str:
    text = str(value or "").strip()

    if text:
        try:
            parsed = datetime.strptime(text, "%Y-%m")
            return parsed.strftime("%Y-%m")
        except ValueError:
            pass

    return date.today().strftime("%Y-%m")



def _payment_period_value(value: str | None) -> str:
    """
    Period 2026-08 means:
      21 Jul 2026 -> 20 Aug 2026.
    """
    raw = str(value or "").strip()

    if raw:
        try:
            parsed = datetime.strptime(raw, "%Y-%m")
            return parsed.strftime("%Y-%m")
        except ValueError:
            pass

    today = date.today()

    if today.day >= 21:
        return today.strftime("%Y-%m")

    previous_month = (
        today.replace(day=1)
        - timedelta(days=1)
    )

    return previous_month.strftime("%Y-%m")


def _payment_period_bounds(
    period: str,
) -> tuple[date, date, date]:

    parsed = datetime.strptime(
        period,
        "%Y-%m",
    ).date()

    period_end = date(
        parsed.year,
        parsed.month,
        20,
    )

    previous_month = (
        period_end.replace(day=1)
        - timedelta(days=1)
    )

    period_start = date(
        previous_month.year,
        previous_month.month,
        21,
    )

    available_from = date(
        parsed.year,
        parsed.month,
        21,
    )

    return (
        period_start,
        period_end,
        available_from,
    )


def _payment_payday(period: str) -> date:
    """
    Return the last Monday-Friday date of the
    payment month.

    V1 intentionally does not adjust for UK
    bank holidays.
    """

    parsed = datetime.strptime(
        period,
        "%Y-%m",
    ).date()

    if parsed.month == 12:
        next_month = date(
            parsed.year + 1,
            1,
            1,
        )
    else:
        next_month = date(
            parsed.year,
            parsed.month + 1,
            1,
        )

    payday = next_month - timedelta(days=1)

    while payday.weekday() >= 5:
        payday -= timedelta(days=1)

    return payday


def _safe_iso_date(value) -> date | None:
    raw = str(value or "").strip()

    if not raw:
        return None

    try:
        return date.fromisoformat(raw)
    except ValueError:
        return None


def _payment_arrangements(session) -> dict[str, str]:

    profiles = {
        item.staff_member_id: item
        for item in session.scalars(
            select(StaffPaymentProfile)
        )
    }

    result = {}

    for member in session.scalars(
        select(StaffMember)
    ):
        profile = profiles.get(member.id)

        arrangement = (
            str(
                profile.payment_arrangement
                if profile
                else "self_employed_contractor"
            )
            .strip()
            or "self_employed_contractor"
        )

        result[
            str(member.staff_source_ref or "")
        ] = arrangement

    return result


def _period_submission_values(
    submission: StaffPayrollSubmission,
    period_start: date,
    period_end: date,
    fallback_arrangement: str,
) -> dict:

    payload = _payload(submission)

    arrangement = (
        str(
            payload.get("payment_arrangement")
            or fallback_arrangement
            or "self_employed_contractor"
        )
        .strip()
    )

    timesheets = []

    for item in payload.get(
        "timesheet_entries",
        [],
    ):
        if not isinstance(item, dict):
            continue

        work_date = _safe_iso_date(
            item.get("work_date")
        )

        if (
            work_date is None
            or work_date < period_start
            or work_date > period_end
        ):
            continue

        paid_minutes = _minutes(
            item.get("start_time"),
            item.get("finish_time"),
            item.get("break_minutes"),
        )

        timesheets.append(
            {
                **item,
                "_work_date": work_date,
                "_paid_minutes": paid_minutes,
            }
        )

    expenses = []

    for item in payload.get(
        "expense_claims",
        [],
    ):
        if not isinstance(item, dict):
            continue

        expense_date = None

        for key in (
            "expense_date",
            "claim_date",
            "work_date",
            "date",
        ):
            expense_date = _safe_iso_date(
                item.get(key)
            )

            if expense_date is not None:
                break

        if expense_date is None:
            expense_date = submission.week_start

        if (
            expense_date < period_start
            or expense_date > period_end
        ):
            continue

        expenses.append(
            {
                **item,
                "_expense_date": expense_date,
                "_amount": float(
                    _money(
                        item.get("amount")
                    )
                ),
            }
        )

    paid_minutes = sum(
        int(item["_paid_minutes"] or 0)
        for item in timesheets
    )

    hours = paid_minutes / 60

    hourly_rate = float(
        _money(
            payload.get("hourly_rate")
        )
    )

    expense_total = sum(
        float(item["_amount"])
        for item in expenses
    )

    if arrangement == "external_payroll":
        gross_pay = 0.0
        total_due = 0.0
    else:
        gross_pay = hours * hourly_rate
        total_due = gross_pay + expense_total

    source_records = []

    for source in payload.get(
        "approved_source_records",
        [],
    ):
        if not isinstance(source, dict):
            continue

        work_date = _safe_iso_date(
            source.get("work_date")
        )

        if (
            work_date is None
            or work_date < period_start
            or work_date > period_end
        ):
            continue

        source_records.append(
            {
                **source,
                "_work_date": work_date,
            }
        )

    return {
        "arrangement": arrangement,
        "timesheets": timesheets,
        "expenses": expenses,
        "source_records": source_records,
        "paid_minutes": paid_minutes,
        "hours_value": hours,
        "hourly_rate_value": hourly_rate,
        "gross_pay_value": gross_pay,
        "expenses_value": expense_total,
        "total_due_value": total_due,
    }


def _payroll_totals(submission: StaffPayrollSubmission) -> dict:
    payload = _payload(submission)
    timesheets = payload["timesheet_entries"]
    expenses = payload["expense_claims"]

    paid_minutes = sum(
        _minutes(
            item.get("start_time"),
            item.get("finish_time"),
            item.get("break_minutes"),
        )
        for item in timesheets
        if isinstance(item, dict)
    )

    expense_total = sum(
        float(_money(item.get("amount")))
        for item in expenses
        if isinstance(item, dict)
    )

    hourly_rate = float(
        _money(payload.get("hourly_rate"))
    )

    hours = paid_minutes / 60
    gross_pay = hours * hourly_rate
    total_due = gross_pay + expense_total

    return {
        "hours_value": hours,
        "hourly_rate_value": hourly_rate,
        "gross_pay_value": gross_pay,
        "expenses_value": expense_total,
        "total_due_value": total_due,
        "hours": f"{hours:.2f}",
        "hourly_rate": f"{hourly_rate:.2f}",
        "gross_pay": f"{gross_pay:.2f}",
        "expenses": f"{expense_total:.2f}",
        "total_due": f"{total_due:.2f}",
    }


def build_payroll_admin_router(templates: Jinja2Templates) -> APIRouter:
    router = APIRouter(prefix="/staff/payroll", tags=["payroll-admin"])

    @router.get(
        "/unlock",
        response_class=HTMLResponse,
        include_in_schema=False,
    )
    async def payroll_unlock(
        request: Request,
        error: str = "",
        sent: str = "",
    ):
        email = _require_staff_super_admin(request)

        if not _payroll_email_allowed(request, email):
            _clear_pending_payroll_otp(request)
            _clear_payroll_unlock(request)
            raise HTTPException(
                403,
                "This Super Admin account is not authorised for Payroll Admin.",
            )

        return templates.TemplateResponse(
            request=request,
            name="payroll_admin_unlock.html",
            context={
                "staff_email": email,
                "error": error,
                "sent": sent == "1",
                "csrf_token": _payroll_csrf(request),
            },
            headers={
                "Cache-Control": "no-store, private",
                "Pragma": "no-cache",
            },
        )

    @router.post(
        "/unlock/send-code",
        include_in_schema=False,
    )
    async def payroll_unlock_send_code(
        request: Request,
        csrf_token: str = Form(...),
    ):
        email = _require_staff_super_admin(request)

        if not _valid_payroll_csrf(request, csrf_token):
            raise HTTPException(
                403,
                "Invalid payroll security request.",
            )

        if not _payroll_email_allowed(request, email):
            _clear_pending_payroll_otp(request)
            _clear_payroll_unlock(request)
            raise HTTPException(
                403,
                "This Super Admin account is not authorised for Payroll Admin.",
            )

        now = time.time()
        last_sent = float(
            request.session.get("payroll_otp_last_sent_at") or 0
        )

        if now - last_sent < PAYROLL_OTP_RESEND_SECONDS:
            return RedirectResponse(
                "/staff/payroll/unlock?error="
                + quote("Please wait before requesting another code."),
                status_code=303,
            )

        settings = request.app.state.settings
        code = f"{secrets.randbelow(1_000_000):06d}"

        try:
            _send_payroll_otp(
                settings,
                email,
                code,
            )
        except Exception:
            _clear_pending_payroll_otp(request)

            return RedirectResponse(
                "/staff/payroll/unlock?error="
                + quote(
                    "The security email could not be sent. "
                    "Please try again shortly."
                ),
                status_code=303,
            )

        request.session["payroll_otp_email"] = email
        request.session["payroll_otp_digest"] = _payroll_otp_digest(
            settings.secret_key,
            code,
        )
        request.session["payroll_otp_expires_at"] = (
            now + int(settings.payroll_otp_ttl_seconds)
        )
        request.session["payroll_otp_attempts"] = 0
        request.session["payroll_otp_last_sent_at"] = now

        return RedirectResponse(
            "/staff/payroll/unlock?sent=1",
            status_code=303,
        )

    @router.post(
        "/unlock/verify",
        include_in_schema=False,
    )
    async def payroll_unlock_verify(
        request: Request,
        code: str = Form(...),
        csrf_token: str = Form(...),
    ):
        email = _require_staff_super_admin(request)

        if not _valid_payroll_csrf(request, csrf_token):
            raise HTTPException(
                403,
                "Invalid payroll security request.",
            )

        if not _payroll_email_allowed(request, email):
            _clear_pending_payroll_otp(request)
            _clear_payroll_unlock(request)
            raise HTTPException(
                403,
                "This Super Admin account is not authorised for Payroll Admin.",
            )

        now = time.time()
        settings = request.app.state.settings

        pending_email = str(
            request.session.get("payroll_otp_email") or ""
        ).lower()

        expected = str(
            request.session.get("payroll_otp_digest") or ""
        )

        expires_at = float(
            request.session.get("payroll_otp_expires_at") or 0
        )

        attempts = int(
            request.session.get("payroll_otp_attempts") or 0
        )

        clean_code = str(code or "").strip()

        invalid = (
            not expected
            or pending_email != email
            or now > expires_at
            or attempts >= PAYROLL_OTP_MAX_ATTEMPTS
            or len(clean_code) != 6
            or not clean_code.isdigit()
        )

        if invalid:
            _clear_pending_payroll_otp(request)

            return RedirectResponse(
                "/staff/payroll/unlock?error="
                + quote(
                    "That security code has expired or is no longer valid."
                ),
                status_code=303,
            )

        supplied = _payroll_otp_digest(
            settings.secret_key,
            clean_code,
        )

        if not hmac.compare_digest(expected, supplied):
            attempts += 1
            request.session["payroll_otp_attempts"] = attempts

            if attempts >= PAYROLL_OTP_MAX_ATTEMPTS:
                _clear_pending_payroll_otp(request)

                message = (
                    "Too many incorrect attempts. "
                    "Request a new security code."
                )
            else:
                message = "That security code was not recognised."

            return RedirectResponse(
                "/staff/payroll/unlock?error=" + quote(message),
                status_code=303,
            )

        _clear_pending_payroll_otp(request)

        request.session["payroll_otp_verified_at"] = now
        request.session["payroll_otp_last_activity"] = now

        return RedirectResponse(
            "/staff/payroll",
            status_code=303,
        )

    @router.get(
        "/staff-payment-setup",
        response_class=HTMLResponse,
        include_in_schema=False,
    )
    async def payroll_staff_payment_setup(
        request: Request,
    ):
        # Uses the existing Payroll Admin protection:
        # Super Admin + authorised Payroll identity +
        # Payroll OTP session.
        _require_super_admin(request)

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        with session_factory() as session:

            members = list(
                session.scalars(
                    select(StaffMember)
                    .where(
                        StaffMember.active.is_(True)
                    )
                    .order_by(
                        StaffMember.staff_name.asc()
                    )
                )
            )

            profiles = {
                item.staff_member_id: item
                for item in session.scalars(
                    select(StaffPaymentProfile)
                )
            }

            all_rates = list(
                session.scalars(
                    select(StaffPayRate)
                    .order_by(
                        StaffPayRate.staff_member_id.asc(),
                        StaffPayRate.effective_from.desc(),
                        StaffPayRate.created_at.desc(),
                    )
                )
            )

            rate_history = {}

            for rate in all_rates:
                rate_history.setdefault(
                    rate.staff_member_id,
                    [],
                ).append(rate)

            rows = []

            for member in members:

                profile = profiles.get(
                    member.id
                )

                rates = rate_history.get(
                    member.id,
                    [],
                )

                current_rate = None

                today = date.today()

                for rate in rates:
                    if rate.effective_from <= today:
                        current_rate = rate
                        break

                arrangement = (
                    str(
                        profile.payment_arrangement
                        if profile
                        else "self_employed_contractor"
                    )
                    .strip()
                )

                rows.append(
                    {
                        "member": member,
                        "profile": profile,
                        "arrangement": arrangement,
                        "arrangement_label": (
                            "External Payroll"
                            if arrangement
                            == "external_payroll"
                            else "Self-Employed Contractor"
                        ),
                        "current_rate": current_rate,
                        "rates": rates,
                    }
                )

        return templates.TemplateResponse(
            request=request,
            name="payroll_admin_staff.html",
            context={
                "rows": rows,
                "csrf_token": _payroll_csrf(
                    request
                ),
                "staff_email":
                    request.session.get(
                        "staff_email",
                        "",
                    ),
            },
            headers={
                "Cache-Control":
                    "no-store, private",
            },
        )


    @router.post(
        "/staff-payment-setup/{staff_id}/arrangement",
        include_in_schema=False,
    )
    async def payroll_staff_arrangement_save(
        request: Request,
        staff_id: str,
        payment_arrangement: str = Form(...),
        notes: str = Form(""),
        csrf_token: str = Form(...),
    ):
        _require_super_admin(request)

        target = "/staff/payroll/staff-payment-setup"

        if not _valid_payroll_csrf(
            request,
            csrf_token,
        ):
            return RedirectResponse(
                target + "?error=csrf",
                status_code=303,
            )

        arrangement = str(
            payment_arrangement or ""
        ).strip()

        allowed = {
            "self_employed_contractor",
            "external_payroll",
        }

        if arrangement not in allowed:
            return RedirectResponse(
                target + "?error=arrangement",
                status_code=303,
            )

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        with session_factory() as session:

            member = session.get(
                StaffMember,
                str(staff_id or "").strip(),
            )

            if member is None:
                return RedirectResponse(
                    target + "?error=staff",
                    status_code=303,
                )

            profile = session.scalar(
                select(StaffPaymentProfile)
                .where(
                    StaffPaymentProfile.staff_member_id
                    == member.id
                )
            )

            if profile is None:
                profile = StaffPaymentProfile(
                    staff_member_id=member.id,
                )
                session.add(profile)

            profile.payment_arrangement = (
                arrangement
            )

            profile.notes = str(
                notes or ""
            ).strip()[:4000]

            profile.updated_by = str(
                request.session.get(
                    "staff_email",
                    "",
                )
            )[:320]

            profile.updated_at = datetime.now()

            session.commit()

        return RedirectResponse(
            target + "?saved=arrangement",
            status_code=303,
        )


    @router.post(
        "/staff-payment-setup/{staff_id}/rate",
        include_in_schema=False,
    )
    async def payroll_staff_rate_save(
        request: Request,
        staff_id: str,
        hourly_rate: str = Form(...),
        effective_from: str = Form(...),
        notes: str = Form(""),
        csrf_token: str = Form(...),
    ):
        _require_super_admin(request)

        target = "/staff/payroll/staff-payment-setup"

        if not _valid_payroll_csrf(
            request,
            csrf_token,
        ):
            return RedirectResponse(
                target + "?error=csrf",
                status_code=303,
            )

        try:
            amount = Decimal(
                str(hourly_rate or "").strip()
            ).quantize(
                Decimal("0.01"),
                rounding=ROUND_HALF_UP,
            )

            effective_date = (
                datetime.strptime(
                    str(
                        effective_from
                        or ""
                    ).strip(),
                    "%Y-%m-%d",
                ).date()
            )

        except (
            InvalidOperation,
            ValueError,
            TypeError,
        ):
            return RedirectResponse(
                target + "?error=rate",
                status_code=303,
            )

        if amount <= 0:
            return RedirectResponse(
                target + "?error=rate",
                status_code=303,
            )

        pence = int(
            (
                amount * Decimal("100")
            ).to_integral_value(
                rounding=ROUND_HALF_UP
            )
        )

        if pence <= 0:
            return RedirectResponse(
                target + "?error=rate",
                status_code=303,
            )

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        with session_factory() as session:

            member = session.get(
                StaffMember,
                str(staff_id or "").strip(),
            )

            if member is None:
                return RedirectResponse(
                    target + "?error=staff",
                    status_code=303,
                )

            # One authoritative rate for a member/effective date.
            # A correction to the same effective date updates that
            # record; a genuinely new rate uses another date and
            # therefore preserves the historic rate.
            existing = session.scalar(
                select(StaffPayRate).where(
                    StaffPayRate.staff_member_id
                    == member.id,
                    StaffPayRate.effective_from
                    == effective_date,
                )
            )

            actor = str(
                request.session.get(
                    "staff_email",
                    "",
                )
            )[:320]

            if existing is not None:
                existing.hourly_rate_pence = (
                    pence
                )
                existing.notes = str(
                    notes or ""
                ).strip()[:4000]
                existing.created_by = actor

            else:
                session.add(
                    StaffPayRate(
                        staff_member_id=member.id,
                        effective_from=effective_date,
                        hourly_rate_pence=pence,
                        notes=str(
                            notes or ""
                        ).strip()[:4000],
                        created_by=actor,
                    )
                )

            session.commit()

        return RedirectResponse(
            target + "?saved=rate",
            status_code=303,
        )


    @router.get("", response_class=HTMLResponse, include_in_schema=False)
    async def payroll_list(request: Request):
        _require_super_admin(request)
        session_factory = getattr(request.app.state, "session_factory", None)
        if session_factory is None:
            raise HTTPException(503, "Database unavailable.")

        with session_factory() as session:
            submissions = list(
                session.scalars(
                    select(StaffPayrollSubmission)
                    .where(
                        StaffPayrollSubmission.status
                        == "submitted",
                        StaffPayrollSubmission.paid_at
                        .is_(None),
                    )
                    .order_by(
                        StaffPayrollSubmission.week_start.desc(),
                        StaffPayrollSubmission.updated_at.desc(),
                    )
                )
            )

            rows = []
            for submission in submissions:
                payload = _payload(submission)
                timesheets = payload["timesheet_entries"]
                expenses = payload["expense_claims"]

                paid_minutes = sum(
                    _minutes(
                        item.get("start_time"),
                        item.get("finish_time"),
                        item.get("break_minutes"),
                    )
                    for item in timesheets
                    if isinstance(item, dict)
                )

                expense_total = sum(
                    float(_money(item.get("amount")))
                    for item in expenses
                    if isinstance(item, dict)
                )

                hourly_rate = float(
                    _money(payload.get("hourly_rate"))
                )
                gross_pay = (paid_minutes / 60) * hourly_rate
                total_due = gross_pay + expense_total

                rows.append(
                    {
                        "id": submission.id,
                        "staff_name": submission.staff_name,
                        "staff_email": submission.staff_email,
                        "week_start": submission.week_start,
                        "status": submission.status,
                        "approved_at": submission.approved_at,
                        "management_state": (
                            "ready-to-pay"
                            if submission.approved_at is not None
                            else "needs-approval"
                        ),
                        "management_label": (
                            "READY TO PAY"
                            if submission.approved_at is not None
                            else "NEEDS APPROVAL"
                        ),
                        "submitted_at": submission.submitted_at,
                        "updated_at": submission.updated_at,
                        "hours": f"{paid_minutes / 60:.2f}",
                        "hourly_rate": f"{hourly_rate:.2f}",
                        "gross_pay": f"{gross_pay:.2f}",
                        "expenses": f"{expense_total:.2f}",
                        "total_due": f"{total_due:.2f}",
                    }
                )

        return templates.TemplateResponse(
            request=request,
            name="payroll_admin_list.html",
            context={
                "rows": rows,
                "staff_email": request.session.get("staff_email", ""),
            },
            headers={"Cache-Control": "no-store, private"},
        )

    @router.get(
        "/monthly/export.xlsx",
        include_in_schema=False,
    )
    async def payroll_monthly_export(
        request: Request,
    ):
        _require_super_admin(request)

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        period = _payment_period_value(
            request.query_params.get("month")
        )

        (
            period_start,
            period_end,
            available_from,
        ) = _payment_period_bounds(period)

        payday = _payment_payday(period)

        if date.today() < available_from:
            raise HTTPException(
                409,
                (
                    "This payment-period report "
                    "will be available from "
                    f"{available_from.strftime('%d %B %Y')}."
                ),
            )

        with session_factory() as session:

            submissions = list(
                session.scalars(
                    select(StaffPayrollSubmission)
                    .order_by(
                        StaffPayrollSubmission.staff_name,
                        StaffPayrollSubmission.week_start,
                    )
                )
            )

            arrangements = _payment_arrangements(
                session
            )

            rows = []

            for submission in submissions:

                fallback = arrangements.get(
                    str(
                        submission.staff_source_ref
                        or ""
                    ),
                    "self_employed_contractor",
                )

                values = _period_submission_values(
                    submission,
                    period_start,
                    period_end,
                    fallback,
                )

                if (
                    not values["timesheets"]
                    and not values["expenses"]
                ):
                    continue

                approved = (
                    submission.status == "submitted"
                    and submission.approved_at
                    is not None
                )

                paid = (
                    approved
                    and submission.paid_at
                    is not None
                )

                rows.append(
                    {
                        "submission": submission,
                        "staff_name":
                            submission.staff_name,
                        "staff_email":
                            submission.staff_email,
                        "staff_source_ref":
                            submission.staff_source_ref,
                        "approved": approved,
                        "paid": paid,
                        **values,
                    }
                )

        workbook = Workbook()

        # ==================================================
        # MANAGEMENT SUMMARY
        # ==================================================

        sheet = workbook.active
        sheet.title = "Management Summary"

        sheet.append(
            ["STUPHIE LTD PAYMENT REPORT"]
        )

        sheet.append(
            [
                "Payment period",
                (
                    f"{period_start.strftime('%d %b %Y')}"
                    " - "
                    f"{period_end.strftime('%d %b %Y')}"
                ),
            ]
        )

        sheet.append(
            [
                "Available from",
                available_from.strftime(
                    "%d %b %Y"
                ),
            ]
        )

        sheet.append(
            [
                "Payday",
                payday.strftime(
                    "%A %d %B %Y"
                ),
            ]
        )

        sheet.append([])

        headings = [
            "Staff member",
            "Email",
            "Payment arrangement",
            "Hours",
            "Agreed hourly rate",
            "Approved work value",
            "Expenses",
            "Outstanding amount due",
            "Approved records",
            "Paid records",
            "Pending records",
        ]

        sheet.append(headings)

        for cell in sheet[6]:
            cell.font = Font(bold=True)

        summary = {}

        for row in rows:

            key = (
                row["staff_source_ref"]
                or row["staff_email"]
                or row["staff_name"]
            )

            item = summary.setdefault(
                key,
                {
                    "staff_name":
                        row["staff_name"],
                    "staff_email":
                        row["staff_email"],
                    "arrangement":
                        row["arrangement"],
                    "hours": 0.0,
                    "rate":
                        row["hourly_rate_value"],
                    "gross": 0.0,
                    "expenses": 0.0,
                    "due": 0.0,
                    "approved": 0,
                    "paid": 0,
                    "pending": 0,
                },
            )

            item["hours"] += (
                row["hours_value"]
            )

            if row["paid"]:
                item["paid"] += 1
            elif row["approved"]:
                item["approved"] += 1
            else:
                item["pending"] += 1

            if (
                row["arrangement"]
                == "self_employed_contractor"
                and row["approved"]
            ):
                item["gross"] += (
                    row["gross_pay_value"]
                )

                item["expenses"] += (
                    row["expenses_value"]
                )

                if not row["paid"]:
                    item["due"] += (
                        row["total_due_value"]
                    )

        for item in sorted(
            summary.values(),
            key=lambda x: (
                x["staff_name"] or ""
            ).lower(),
        ):

            external = (
                item["arrangement"]
                == "external_payroll"
            )

            sheet.append(
                [
                    item["staff_name"],
                    item["staff_email"],
                    (
                        "External Payroll"
                        if external
                        else "Self-Employed Contractor"
                    ),
                    item["hours"],
                    (
                        None
                        if external
                        else item["rate"]
                    ),
                    (
                        None
                        if external
                        else item["gross"]
                    ),
                    item["expenses"],
                    (
                        None
                        if external
                        else item["due"]
                    ),
                    item["approved"],
                    item["paid"],
                    item["pending"],
                ]
            )

        for column in (
            "E",
            "F",
            "G",
            "H",
        ):
            for cell in sheet[column][6:]:
                cell.number_format = (
                    '£#,##0.00'
                )

        sheet.freeze_panes = "A7"

        # ==================================================
        # CONTRACTOR DETAIL
        # ==================================================

        contractor = workbook.create_sheet(
            "Contractor Detail"
        )

        contractor.append(
            [
                "Staff member",
                "Email",
                "Payroll week starting",
                "Hours in period",
                "Hourly rate",
                "Work value",
                "Expenses",
                "Outstanding amount due",
                "Management status",
                "Payment status",
                "Payment reference",
            ]
        )

        for cell in contractor[1]:
            cell.font = Font(bold=True)

        for row in rows:

            if (
                row["arrangement"]
                != "self_employed_contractor"
            ):
                continue

            submission = row["submission"]

            contractor.append(
                [
                    row["staff_name"],
                    row["staff_email"],
                    submission.week_start,
                    row["hours_value"],
                    row["hourly_rate_value"],
                    (
                        row["gross_pay_value"]
                        if row["approved"]
                        else 0.0
                    ),
                    (
                        row["expenses_value"]
                        if row["approved"]
                        else 0.0
                    ),
                    (
                        row["total_due_value"]
                        if (
                            row["approved"]
                            and not row["paid"]
                        )
                        else 0.0
                    ),
                    (
                        "APPROVED"
                        if row["approved"]
                        else "PENDING"
                    ),
                    (
                        "PAID"
                        if row["paid"]
                        else (
                            "READY TO PAY"
                            if row["approved"]
                            else "NOT YET PAYABLE"
                        )
                    ),
                    (
                        submission.payment_reference
                        or ""
                    ),
                ]
            )

        # ==================================================
        # EXTERNAL PAYROLL
        # ==================================================

        external = workbook.create_sheet(
            "External Payroll"
        )

        external.append(
            [
                "Staff member",
                "Email",
                "Payroll week starting",
                "Hours in period",
                "Management status",
            ]
        )

        for cell in external[1]:
            cell.font = Font(bold=True)

        for row in rows:

            if (
                row["arrangement"]
                != "external_payroll"
            ):
                continue

            external.append(
                [
                    row["staff_name"],
                    row["staff_email"],
                    row["submission"].week_start,
                    row["hours_value"],
                    (
                        "APPROVED"
                        if row["approved"]
                        else "PENDING"
                    ),
                ]
            )

        # ==================================================
        # WORK AUDIT
        # ==================================================

        audit = workbook.create_sheet(
            "Work Audit"
        )

        audit.append(
            [
                "Staff member",
                "Payment arrangement",
                "Work date",
                "Start",
                "Finish",
                "Break minutes",
                "Paid hours",
                "Notes",
                "Payroll week starting",
                "Management status",
            ]
        )

        for cell in audit[1]:
            cell.font = Font(bold=True)

        for row in rows:

            arrangement_label = (
                "External Payroll"
                if row["arrangement"]
                == "external_payroll"
                else "Self-Employed Contractor"
            )

            for item in row["timesheets"]:

                audit.append(
                    [
                        row["staff_name"],
                        arrangement_label,
                        item["_work_date"],
                        item.get(
                            "start_time",
                            "",
                        ),
                        item.get(
                            "finish_time",
                            "",
                        ),
                        int(
                            item.get(
                                "break_minutes",
                                0,
                            )
                            or 0
                        ),
                        (
                            int(
                                item[
                                    "_paid_minutes"
                                ]
                                or 0
                            )
                            / 60
                        ),
                        str(
                            item.get(
                                "notes",
                                "",
                            )
                            or ""
                        ),
                        row[
                            "submission"
                        ].week_start,
                        (
                            "APPROVED"
                            if row["approved"]
                            else "PENDING"
                        ),
                    ]
                )

        # ==================================================
        # SOURCE AUDIT
        # ==================================================

        source = workbook.create_sheet(
            "Approved Source Audit"
        )

        source.append(
            [
                "Staff member",
                "Work date",
                "Source type",
                "Source ID",
                "Event",
                "Work type",
                "Description",
                "Paid minutes",
            ]
        )

        for cell in source[1]:
            cell.font = Font(bold=True)

        for row in rows:

            for record in row[
                "source_records"
            ]:

                source.append(
                    [
                        row["staff_name"],
                        record["_work_date"],
                        str(
                            record.get(
                                "source_type",
                                "",
                            )
                            or ""
                        ),
                        str(
                            record.get(
                                "source_id",
                                "",
                            )
                            or ""
                        ),
                        str(
                            record.get(
                                "event_name",
                                "",
                            )
                            or ""
                        ),
                        str(
                            record.get(
                                "work_type",
                                "",
                            )
                            or ""
                        ),
                        str(
                            record.get(
                                "description",
                                "",
                            )
                            or ""
                        ),
                        int(
                            record.get(
                                "paid_minutes",
                                0,
                            )
                            or 0
                        ),
                    ]
                )

        # Friendly widths
        for worksheet in (
            sheet,
            contractor,
            external,
            audit,
            source,
        ):
            for column_cells in (
                worksheet.columns
            ):
                width = max(
                    (
                        len(
                            str(
                                cell.value
                                if cell.value
                                is not None
                                else ""
                            )
                        )
                        for cell
                        in column_cells
                    ),
                    default=10,
                )

                worksheet.column_dimensions[
                    column_cells[
                        0
                    ].column_letter
                ].width = min(
                    max(width + 2, 12),
                    45,
                )

        output = BytesIO()

        workbook.save(output)
        output.seek(0)

        filename = (
            "stuphie-payment-report-"
            f"{period}-21-to-20.xlsx"
        )

        return StreamingResponse(
            output,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                    (
                        "attachment; "
                        f'filename="{filename}"'
                    ),
                "Cache-Control":
                    "no-store, private",
            },
        )


    @router.get(
        "/monthly",
        response_class=HTMLResponse,
        include_in_schema=False,
    )
    async def payroll_monthly(
        request: Request,
    ):
        _require_super_admin(request)

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        period = _payment_period_value(
            request.query_params.get("month")
        )

        (
            period_start,
            period_end,
            available_from,
        ) = _payment_period_bounds(period)

        report_available = (
            date.today() >= available_from
        )

        payday = _payment_payday(period)

        with session_factory() as session:

            submissions = list(
                session.scalars(
                    select(
                        StaffPayrollSubmission
                    )
                    .order_by(
                        StaffPayrollSubmission.staff_name,
                        StaffPayrollSubmission.week_start,
                    )
                )
            )

            arrangements = (
                _payment_arrangements(
                    session
                )
            )

            staff = {}

            for submission in submissions:

                fallback = arrangements.get(
                    str(
                        submission.staff_source_ref
                        or ""
                    ),
                    "self_employed_contractor",
                )

                values = (
                    _period_submission_values(
                        submission,
                        period_start,
                        period_end,
                        fallback,
                    )
                )

                if (
                    not values["timesheets"]
                    and not values["expenses"]
                ):
                    continue

                key = (
                    submission.staff_source_ref
                    or submission.staff_email
                    or submission.staff_name
                )

                item = staff.setdefault(
                    key,
                    {
                        "staff_name":
                            submission.staff_name,
                        "staff_email":
                            submission.staff_email,
                        "arrangement":
                            values["arrangement"],
                        "hours_value": 0.0,
                        "gross_pay_value": 0.0,
                        "expenses_value": 0.0,
                        "total_due_value": 0.0,
                        "approved_records": 0,
                        "pending_records": 0,
                        "paid_records": 0,
                    },
                )

                item["hours_value"] += (
                    values["hours_value"]
                )

                approved = (
                    submission.status == "submitted"
                    and submission.approved_at
                    is not None
                )

                paid = (
                    approved
                    and submission.paid_at
                    is not None
                )

                if paid:
                    item["paid_records"] += 1

                elif approved:
                    item["approved_records"] += 1

                else:
                    item["pending_records"] += 1

                if (
                    values["arrangement"]
                    == "self_employed_contractor"
                    and approved
                ):
                    item[
                        "gross_pay_value"
                    ] += values[
                        "gross_pay_value"
                    ]

                    item[
                        "expenses_value"
                    ] += values[
                        "expenses_value"
                    ]

                    if not paid:
                        item[
                            "total_due_value"
                        ] += values[
                            "total_due_value"
                        ]

                elif (
                    values["arrangement"]
                    == "external_payroll"
                ):
                    item[
                        "expenses_value"
                    ] += values[
                        "expenses_value"
                    ]

            rows = sorted(
                staff.values(),
                key=lambda item: (
                    item["staff_name"]
                    or ""
                ).lower(),
            )

            contractor_rows = [
                row
                for row in rows
                if row["arrangement"]
                == "self_employed_contractor"
            ]

            external_rows = [
                row
                for row in rows
                if row["arrangement"]
                == "external_payroll"
            ]

            totals = {
                "contractor_hours":
                    sum(
                        row["hours_value"]
                        for row
                        in contractor_rows
                    ),
                "external_hours":
                    sum(
                        row["hours_value"]
                        for row
                        in external_rows
                    ),
                "contractor_due":
                    sum(
                        row["total_due_value"]
                        for row
                        in contractor_rows
                    ),
                "expenses":
                    sum(
                        row["expenses_value"]
                        for row
                        in rows
                    ),
                "staff_count":
                    len(rows),
            }

            for item in rows:

                item["hours"] = (
                    f'{item["hours_value"]:.2f}'
                )

                item["gross_pay"] = (
                    f'{item["gross_pay_value"]:.2f}'
                )

                item["expenses"] = (
                    f'{item["expenses_value"]:.2f}'
                )

                item["total_due"] = (
                    f'{item["total_due_value"]:.2f}'
                )

        return templates.TemplateResponse(
            request=request,
            name="payroll_admin_monthly.html",
            context={
                "rows": rows,
                "contractor_rows":
                    contractor_rows,
                "external_rows":
                    external_rows,
                "totals": totals,
                "month": period,
                "period_start":
                    period_start,
                "period_end":
                    period_end,
                "available_from":
                    available_from,
                "report_available":
                    report_available,
                "payday":
                    payday,
                "staff_email":
                    request.session.get(
                        "staff_email",
                        "",
                    ),
            },
            headers={
                "Cache-Control":
                    "no-store, private"
            },
        )


    @router.get(
        "/history",
        response_class=HTMLResponse,
        include_in_schema=False,
    )
    async def payroll_history(request: Request):
        _require_super_admin(request)

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        with session_factory() as session:
            submissions = list(
                session.scalars(
                    select(StaffPayrollSubmission)
                    .where(
                        StaffPayrollSubmission.status
                        == "submitted",
                        StaffPayrollSubmission.paid_at
                        .is_not(None),
                    )
                    .order_by(
                        StaffPayrollSubmission.paid_at.desc(),
                        StaffPayrollSubmission.week_start.desc(),
                    )
                )
            )

            rows = []

            for submission in submissions:
                totals = _payroll_totals(submission)

                rows.append(
                    {
                        "id": submission.id,
                        "staff_name": submission.staff_name,
                        "staff_email": submission.staff_email,
                        "week_start": submission.week_start,
                        "hours": totals["hours"],
                        "hourly_rate": totals["hourly_rate"],
                        "gross_pay": totals["gross_pay"],
                        "expenses": totals["expenses"],
                        "total_due": totals["total_due"],
                        "submitted_at": submission.submitted_at,
                        "approved_at": submission.approved_at,
                        "approved_by": submission.approved_by,
                        "paid_at": submission.paid_at,
                        "paid_by": submission.paid_by,
                        "payment_reference":
                            submission.payment_reference,
                    }
                )

        return templates.TemplateResponse(
            request=request,
            name="payroll_admin_history.html",
            context={
                "rows": rows,
                "staff_email": request.session.get(
                    "staff_email",
                    "",
                ),
            },
            headers={
                "Cache-Control": "no-store, private",
            },
        )


    @router.get(
        "/{submission_id}",
        response_class=HTMLResponse,
        include_in_schema=False,
    )
    async def payroll_detail(request: Request, submission_id: str):
        _require_super_admin(request)
        session_factory = getattr(request.app.state, "session_factory", None)
        if session_factory is None:
            raise HTTPException(503, "Database unavailable.")

        with session_factory() as session:
            submission = session.get(StaffPayrollSubmission, submission_id)
            if submission is None:
                raise HTTPException(404)

            payload = _payload(submission)
            timesheets = [
                item for item in payload["timesheet_entries"]
                if isinstance(item, dict)
            ]
            expenses = [
                item for item in payload["expense_claims"]
                if isinstance(item, dict)
            ]

            for item in timesheets:
                item["_paid_minutes"] = _minutes(
                    item.get("start_time"),
                    item.get("finish_time"),
                    item.get("break_minutes"),
                )

            for index, item in enumerate(expenses):
                location = _expense_receipt(item)
                item["_receipt_location"] = location
                item["_receipt_filename"] = _expense_filename(item, location)
                item["_index"] = index
                item["_amount"] = _money(item.get("amount"))

            paid_minutes = sum(item["_paid_minutes"] for item in timesheets)
            expense_total = sum(float(item["_amount"]) for item in expenses)

            hourly_rate = float(
                _money(payload.get("hourly_rate"))
            )
            gross_pay = (paid_minutes / 60) * hourly_rate
            total_due = gross_pay + expense_total

            payment_arrangement = str(
                payload.get(
                    "payment_arrangement",
                    "",
                )
            ).strip()

            statement_audit = payload.get(
                "contractor_statement"
            )

            if not isinstance(
                statement_audit,
                dict,
            ):
                statement_audit = {}

            view = {
                "id": submission.id,
                "source_ref": submission.source_ref,
                "staff_source_ref": submission.staff_source_ref,
                "staff_name": submission.staff_name,
                "staff_email": submission.staff_email,
                "week_start": submission.week_start,
                "status": submission.status,
                "emailed_at": submission.emailed_at,
                "submitted_at": submission.submitted_at,
                "approved_at": submission.approved_at,
                "approved_by": submission.approved_by,
                "paid_at": submission.paid_at,
                "paid_by": submission.paid_by,
                "payment_reference": submission.payment_reference,
                "created_at": submission.created_at,
                "updated_at": submission.updated_at,
            }

        return templates.TemplateResponse(
            request=request,
            name="payroll_admin_detail.html",
            context={
                "submission": view,
                "timesheets": timesheets,
                "expenses": expenses,
                "total_hours": f"{paid_minutes / 60:.2f}",
                "hourly_rate": f"{hourly_rate:.2f}",
                "gross_pay": f"{gross_pay:.2f}",
                "total_expenses": f"{expense_total:.2f}",
                "total_due": f"{total_due:.2f}",
                "payment_arrangement":
                    payment_arrangement,
                "statement_audit":
                    statement_audit,
                "staff_email": request.session.get("staff_email", ""),
                "csrf_token": _payroll_csrf(request),
            },
            headers={"Cache-Control": "no-store, private"},
        )

    @router.get(
        "/{submission_id}/statement.pdf",
        include_in_schema=False,
    )
    async def payroll_statement_pdf(
        request: Request,
        submission_id: str,
    ):
        _require_super_admin(request)

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        with session_factory() as session:

            submission = session.get(
                StaffPayrollSubmission,
                submission_id,
            )

            if submission is None:
                raise HTTPException(404)

            if submission.approved_at is None:
                raise HTTPException(
                    409,
                    "Statement is available after management approval.",
                )

            payload = _payload(
                submission
            )

            if (
                str(
                    payload.get(
                        "payment_arrangement",
                        "",
                    )
                ).strip()
                != "self_employed_contractor"
            ):
                raise HTTPException(
                    409,
                    "This payroll record is not a contractor payment.",
                )

            pdf_bytes = (
                _contractor_statement_pdf(
                    request=request,
                    submission=submission,
                    payload=payload,
                )
            )

        filename = (
            "sophies-photography-"
            "contractor-payment-statement-"
            f"{submission.week_start.isoformat()}.pdf"
        )

        return StreamingResponse(
            BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition":
                    (
                        "attachment; "
                        f'filename="{filename}"'
                    ),
                "Cache-Control":
                    "no-store, private",
                "X-Content-Type-Options":
                    "nosniff",
            },
        )


    @router.post(
        "/{submission_id}/send-statement",
        include_in_schema=False,
    )
    async def payroll_send_statement(
        request: Request,
        submission_id: str,
        csrf_token: str = Form(...),
    ):
        sender_email = (
            _require_staff_super_admin(
                request
            )
        )

        _require_super_admin(request)

        if not _valid_payroll_csrf(
            request,
            csrf_token,
        ):
            raise HTTPException(
                403,
                "Invalid statement email request.",
            )

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        settings = (
            request.app.state.settings
        )

        with session_factory() as session:

            submission = session.get(
                StaffPayrollSubmission,
                submission_id,
            )

            if submission is None:
                raise HTTPException(404)

            if submission.approved_at is None:
                raise HTTPException(
                    409,
                    "Statement cannot be sent before management approval.",
                )

            payload = _payload(
                submission
            )

            if (
                str(
                    payload.get(
                        "payment_arrangement",
                        "",
                    )
                ).strip()
                != "self_employed_contractor"
            ):
                raise HTTPException(
                    409,
                    "External-payroll staff do not receive contractor statements.",
                )

            recipient = str(
                submission.staff_email
                or ""
            ).strip()

            if not recipient:
                raise HTTPException(
                    409,
                    "Contractor email address is missing.",
                )

            pdf_bytes = (
                _contractor_statement_pdf(
                    request=request,
                    submission=submission,
                    payload=payload,
                )
            )

            message = EmailMessage()

            message["From"] = (
                f"{settings.smtp_from_name} "
                f"<{settings.smtp_from_email}>"
            )

            message["To"] = recipient

            message["Subject"] = (
                "Your Sophie’s Photography "
                "contractor payment statement"
            )

            display_name = (
                submission.staff_name
                or "there"
            )

            message.set_content(
                (
                    f"Hi {display_name},\n\n"
                    "Attached is your contractor payment "
                    "statement from Sophie’s Photography.\n\n"
                    "This is a record of the approved work, "
                    "hours, expenses and payment value for "
                    "the period shown. It is not a PAYE "
                    "payslip.\n\n"
                    "Kind regards,\n"
                    "Sophie’s Photography"
                )
            )

            message.add_alternative(
                branded_email_html(
                    "Contractor payment statement",
                    (
                        f"Hi {display_name},\n\n"
                        "Attached is your contractor payment "
                        "statement from Sophie’s Photography.\n\n"
                        "It records your approved work, hours, "
                        "expenses and payment value for the "
                        "period shown.\n\n"
                        "This is a self-employed contractor "
                        "payment statement and is not a PAYE "
                        "payslip."
                    ),
                    eyebrow=(
                        "SOPHIE’S PHOTOGRAPHY · "
                        "PAYMENT STATEMENT"
                    ),
                    include_marketing_controls=False,
                ),
                subtype="html",
            )

            attach_sophies_logo(
                message
            )

            filename = (
                "sophies-photography-"
                "contractor-payment-statement-"
                f"{submission.week_start.isoformat()}.pdf"
            )

            message.add_attachment(
                pdf_bytes,
                maintype="application",
                subtype="pdf",
                filename=filename,
            )

            if not settings.smtp_host:
                raise HTTPException(
                    503,
                    "Email service is unavailable.",
                )

            with smtplib.SMTP(
                settings.smtp_host,
                settings.smtp_port,
                timeout=20,
            ) as smtp:

                if settings.smtp_use_tls:
                    smtp.starttls()

                if settings.smtp_username:
                    smtp.login(
                        settings.smtp_username,
                        settings.smtp_password,
                    )

                smtp.send_message(
                    message
                )

            now = (
                datetime.now()
                .astimezone()
            )

            audit = payload.get(
                "contractor_statement"
            )

            if not isinstance(
                audit,
                dict,
            ):
                audit = {}

            audit.update(
                {
                    "last_generated_at":
                        now.isoformat(),
                    "last_sent_at":
                        now.isoformat(),
                    "last_sent_by":
                        sender_email,
                    "sent_to":
                        recipient,
                    "statement_type":
                        "self_employed_contractor_payment_statement",
                }
            )

            payload[
                "contractor_statement"
            ] = audit

            submission.payload_json = (
                json.dumps(
                    payload,
                    separators=(
                        ",",
                        ":",
                    ),
                )
            )

            session.commit()

        request.session.pop(
            "payroll_admin_csrf_token",
            None,
        )

        return RedirectResponse(
            (
                f"/staff/payroll/{submission_id}"
                "?statement_sent=1"
            ),
            status_code=303,
        )


    @router.post(
        "/{submission_id}/approve",
        include_in_schema=False,
    )
    async def payroll_approve(
        request: Request,
        submission_id: str,
        csrf_token: str = Form(...),
    ):
        approver_email = _require_staff_super_admin(request)

        # Approval requires the complete Payroll Admin security
        # boundary, including a recent Payroll OTP unlock.
        _require_super_admin(request)

        if not _valid_payroll_csrf(
            request,
            csrf_token,
        ):
            raise HTTPException(
                403,
                "Invalid payroll approval request.",
            )

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        with session_factory() as session:
            submission = session.get(
                StaffPayrollSubmission,
                submission_id,
            )

            if submission is None:
                raise HTTPException(404)

            if submission.status != "submitted":
                raise HTTPException(
                    409,
                    "Only submitted payroll can be approved.",
                )

            if submission.approved_at is not None:
                raise HTTPException(
                    409,
                    "This payroll has already been approved.",
                )

            submission.approved_at = datetime.now().astimezone()
            submission.approved_by = approver_email

            session.commit()

        # Consume the form token so a browser refresh cannot
        # replay the same approval request.
        request.session.pop(
            "payroll_admin_csrf_token",
            None,
        )

        return RedirectResponse(
            f"/staff/payroll/{submission_id}?approved=1",
            status_code=303,
        )


    @router.post(
        "/{submission_id}/mark-paid",
        include_in_schema=False,
    )
    async def payroll_mark_paid(
        request: Request,
        submission_id: str,
        csrf_token: str = Form(...),
        payment_reference: str = Form(...),
    ):
        payer_email = _require_staff_super_admin(request)

        # Payroll payment sits behind the complete Payroll
        # Admin security boundary, including recent OTP unlock.
        _require_super_admin(request)

        if not _valid_payroll_csrf(
            request,
            csrf_token,
        ):
            raise HTTPException(
                403,
                "Invalid payroll payment request.",
            )

        payment_reference = payment_reference.strip()

        if not payment_reference:
            raise HTTPException(
                400,
                "A payment reference is required.",
            )

        if len(payment_reference) > 200:
            raise HTTPException(
                400,
                "Payment reference is too long.",
            )

        session_factory = getattr(
            request.app.state,
            "session_factory",
            None,
        )

        if session_factory is None:
            raise HTTPException(
                503,
                "Database unavailable.",
            )

        with session_factory() as session:
            submission = session.get(
                StaffPayrollSubmission,
                submission_id,
            )

            if submission is None:
                raise HTTPException(404)

            if submission.status != "submitted":
                raise HTTPException(
                    409,
                    "Only submitted payroll can be paid.",
                )

            if submission.approved_at is None:
                raise HTTPException(
                    409,
                    "Payroll must be approved before payment.",
                )

            if submission.paid_at is not None:
                raise HTTPException(
                    409,
                    "This payroll has already been paid.",
                )

            submission.paid_at = datetime.now().astimezone()
            submission.paid_by = payer_email
            submission.payment_reference = payment_reference

            session.commit()

        # Consume the token so browser refresh cannot replay
        # the same payroll payment request.
        request.session.pop(
            "payroll_admin_csrf_token",
            None,
        )

        return RedirectResponse(
            f"/staff/payroll/{submission_id}?paid=1",
            status_code=303,
        )


    @router.get(
        "/{submission_id}/receipt/{expense_index}",
        include_in_schema=False,
    )
    async def payroll_receipt(
        request: Request,
        submission_id: str,
        expense_index: int,
    ):
        _require_super_admin(request)
        session_factory = getattr(request.app.state, "session_factory", None)
        if session_factory is None:
            raise HTTPException(503, "Database unavailable.")

        with session_factory() as session:
            submission = session.get(StaffPayrollSubmission, submission_id)
            if submission is None:
                raise HTTPException(404)

            expenses = _payload(submission)["expense_claims"]

            if expense_index < 0 or expense_index >= len(expenses):
                raise HTTPException(404)

            expense = expenses[expense_index]
            if not isinstance(expense, dict):
                raise HTTPException(404)

            location = _expense_receipt(expense)
            if not location:
                raise HTTPException(404)

            filename = _expense_filename(expense, location)

        try:
            suffix = Path(filename).suffix
            temporary = download_to_temp(
                request.app.state.settings,
                location,
                suffix=suffix,
            )
        except Exception as exc:
            raise HTTPException(404, "Receipt is unavailable.") from exc

        media_type = (
            mimetypes.guess_type(filename)[0]
            or "application/octet-stream"
        )

        return FileResponse(
            temporary,
            media_type=media_type,
            filename=filename,
            content_disposition_type="inline",
            headers={
                "Cache-Control": "no-store, private",
                "X-Content-Type-Options": "nosniff",
            },
        )

    return router
